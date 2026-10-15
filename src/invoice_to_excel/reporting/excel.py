"""Excel workbook: invoices, VAT breakdown, modelo 303 and 347 summaries, file issues."""

from __future__ import annotations

import io
import os
import tempfile
from decimal import Decimal
from pathlib import Path

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from invoice_to_excel.models import Invoice, Result
from invoice_to_excel.reporting.summaries import (
    SUPPLIER_347_THRESHOLD,
    quarterly_vat_summary,
    supplier_annual_totals,
)
from invoice_to_excel.reporting.tables import (
    INVOICE_COLUMNS,
    extraction_label,
    format_rate,
    notes_text,
    rate_label,
    sorted_invoices,
    status_label,
)

EUR = '#,##0.00" €";[Red]-#,##0.00" €"'
DATE_FORMAT = "yyyy-mm-dd"

_HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
_HEADER_FONT = Font(bold=True, color="FFFFFF")
_BOLD = Font(bold=True)
_TOTAL_BORDER = Border(top=Side(style="thin"))
_REVIEW_FILL = PatternFill("solid", fgColor="FFC7CE")
_LOW_CONFIDENCE_FILL = PatternFill("solid", fgColor="FFD966")
_FLAG_FILL = PatternFill("solid", fgColor="FFEB9C")

_MONEY_COLUMNS = (5, 7, 8, 9, 10)


def _f(value: Decimal | None) -> float | None:
    return float(value) if value is not None else None


def _header(ws: Worksheet, row_no: int, titles: list[str]) -> None:
    for col, title in enumerate(titles, 1):
        cell = ws.cell(row=row_no, column=col, value=title)
        cell.fill, cell.font = _HEADER_FILL, _HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[row_no].height = 24


def _widths(ws: Worksheet, widths: list[int]) -> None:
    for i, width in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = width


def _invoices_sheet(wb: Workbook, invoices: list[Invoice], threshold: float) -> None:
    ws = wb.active
    if ws is None:
        raise RuntimeError("workbook has no active sheet")
    ws.title = "Invoices"
    _header(ws, 1, [c.title for c in INVOICE_COLUMNS])
    _widths(ws, [c.width for c in INVOICE_COLUMNS])
    rows = sorted_invoices(invoices)
    for r, invoice in enumerate(rows, 2):
        values = [
            invoice.supplier,
            invoice.tax_id,
            invoice.number,
            invoice.invoice_date,
            _f(invoice.net_amount),
            rate_label(invoice) or None,
            _f(invoice.vat_amount),
            _f(invoice.surcharge_amount),
            _f(invoice.irpf),
            _f(invoice.total),
            invoice.overall_confidence,
            extraction_label(invoice),
            invoice.filename,
            status_label(invoice),
            notes_text(invoice),
        ]
        for c, value in enumerate(values, 1):
            cell = ws.cell(row=r, column=c, value=value)
            if not invoice.is_ok:
                cell.fill = _REVIEW_FILL
        ws.cell(row=r, column=4).number_format = DATE_FORMAT
        for c in _MONEY_COLUMNS:
            ws.cell(row=r, column=c).number_format = EUR
        ws.cell(row=r, column=11).number_format = "0%"
        ws.cell(row=r, column=14).alignment = Alignment(horizontal="center")
        _highlight_low_confidence(ws, r, invoice, threshold)

    last_row = len(rows) + 1
    if rows:
        ws.auto_filter.ref = f"A1:{get_column_letter(len(INVOICE_COLUMNS))}{last_row}"
        total_row = last_row + 2  # kept apart from the filterable range
        ws.cell(row=total_row, column=1, value="TOTAL (visible rows)").font = _BOLD
        for c in _MONEY_COLUMNS:
            col = get_column_letter(c)
            cell = ws.cell(row=total_row, column=c, value=f"=SUBTOTAL(109,{col}2:{col}{last_row})")
            cell.number_format, cell.font = EUR, _BOLD
        for c in range(1, 11):
            ws.cell(row=total_row, column=c).border = _TOTAL_BORDER
        legend = total_row + 2
        ws.cell(row=legend, column=1, value="Red row: needs review").fill = _REVIEW_FILL
        ws.cell(
            row=legend + 1,
            column=1,
            value=f"Amber cell: field confidence below {threshold:.0%}",
        ).fill = _LOW_CONFIDENCE_FILL
    ws.freeze_panes = "A2"


def _highlight_low_confidence(ws: Worksheet, row: int, invoice: Invoice, threshold: float) -> None:
    for c, column in enumerate(INVOICE_COLUMNS, 1):
        if column.field is None:
            continue
        if column.field == "irpf" and invoice.irpf is None:
            continue
        score = invoice.confidence.get(column.field, 0.0)
        if score < threshold:
            cell = ws.cell(row=row, column=c)
            cell.fill = _LOW_CONFIDENCE_FILL
            cell.comment = Comment(f"Confidence {score:.0%}", "invoice-to-excel")


def _tax_lines_sheet(wb: Workbook, invoices: list[Invoice]) -> None:
    ws = wb.create_sheet("Tax lines")
    titles = [
        "Supplier", "Invoice no.", "Date", "VAT %", "Net amount", "VAT amount",
        "Surcharge %", "Surcharge", "Source file",
    ]  # fmt: skip
    _header(ws, 1, titles)
    _widths(ws, [34, 16, 12, 10, 16, 14, 12, 14, 28])
    r = 2
    for invoice in sorted_invoices(invoices):
        for line in invoice.tax_lines:
            values = [
                invoice.supplier,
                invoice.number,
                invoice.invoice_date,
                format_rate(line.rate),
                _f(line.base),
                _f(line.vat),
                None if line.surcharge_rate is None else format_rate(line.surcharge_rate),
                _f(line.surcharge),
                invoice.filename,
            ]
            for c, value in enumerate(values, 1):
                ws.cell(row=r, column=c, value=value)
            ws.cell(row=r, column=3).number_format = DATE_FORMAT
            for c in (5, 6, 8):
                ws.cell(row=r, column=c).number_format = EUR
            r += 1
    ws.freeze_panes = "A2"
    if r > 2:
        ws.auto_filter.ref = f"A1:{get_column_letter(len(titles))}{r - 1}"


def _vat_sheet(wb: Workbook, invoices: list[Invoice]) -> None:
    ws = wb.create_sheet("VAT 303")
    titles = ["Quarter", "VAT %", "Invoices", "To review", "Net amount", "VAT amount", "Surcharge"]
    _header(ws, 1, titles)
    _widths(ws, [14, 10, 10, 11, 16, 16, 14])
    rows = quarterly_vat_summary(invoices)
    r = 2
    grand = [Decimal(0)] * 3
    for period in dict.fromkeys(row.period for row in rows):
        subtotal = [Decimal(0)] * 3
        for row in (x for x in rows if x.period == period):
            ws.cell(row=r, column=1, value=row.period)
            ws.cell(row=r, column=2, value=format_rate(row.rate))
            ws.cell(row=r, column=3, value=row.invoices)
            ws.cell(row=r, column=4, value=row.to_review)
            for c, amount in enumerate((row.base, row.vat, row.surcharge), 5):
                ws.cell(row=r, column=c, value=_f(amount)).number_format = EUR
            for i, amount in enumerate((row.base, row.vat, row.surcharge)):
                subtotal[i] += amount
            r += 1
        ws.cell(row=r, column=1, value=f"{period} total")
        for c, amount in enumerate(subtotal, 5):
            ws.cell(row=r, column=c, value=_f(amount)).number_format = EUR
        for c in range(1, len(titles) + 1):
            ws.cell(row=r, column=c).font = _BOLD
            ws.cell(row=r, column=c).border = _TOTAL_BORDER
        grand = [g + s for g, s in zip(grand, subtotal, strict=True)]
        r += 2
    if rows:
        ws.cell(row=r, column=1, value="ALL")
        for c, amount in enumerate(grand, 5):
            ws.cell(row=r, column=c, value=_f(amount)).number_format = EUR
        for c in range(1, len(titles) + 1):
            ws.cell(row=r, column=c).font = _BOLD
        r += 2
    ws.cell(
        row=r,
        column=1,
        value="Input VAT by invoice date. Duplicates excluded; invoices to review included.",
    )
    ws.freeze_panes = "A2"


def _suppliers_sheet(wb: Workbook, invoices: list[Invoice]) -> None:
    ws = wb.create_sheet("Suppliers 347")
    titles = ["Year", "Supplier", "Tax ID", "Invoices", "Total incl. VAT", "Over threshold"]
    _header(ws, 1, titles)
    _widths(ws, [10, 38, 14, 10, 18, 16])
    rows = supplier_annual_totals(invoices)
    for r, row in enumerate(rows, 2):
        values = [
            row.year if row.year is not None else "No date",
            row.supplier,
            row.tax_id,
            row.invoices,
            _f(row.gross),
            "YES" if row.over_threshold else None,
        ]
        for c, value in enumerate(values, 1):
            cell = ws.cell(row=r, column=c, value=value)
            if row.over_threshold:
                cell.fill = _FLAG_FILL
        ws.cell(row=r, column=5).number_format = EUR
        ws.cell(row=r, column=6).alignment = Alignment(horizontal="center")
    note = len(rows) + 3
    ws.cell(
        row=note,
        column=1,
        value=(
            f"Modelo 347 threshold: {SUPPLIER_347_THRESHOLD:,.2f} EUR per supplier and year, "
            "VAT included, IRPF not deducted. Duplicates are excluded."
        ),
    )
    ws.freeze_panes = "A2"


def _issues_sheet(wb: Workbook, result: Result) -> None:
    ws = wb.create_sheet("Issues")
    _header(ws, 1, ["File", "Reason"])
    _widths(ws, [32, 110])
    for r, issue in enumerate(result.file_issues, 2):
        ws.cell(row=r, column=1, value=issue.filename)
        ws.cell(row=r, column=2, value=issue.reason)
    if not result.file_issues:
        ws.cell(row=2, column=1, value="No issues: every PDF could be read.")
    ws.freeze_panes = "A2"


def build_workbook(result: Result, confidence_threshold: float = 0.7) -> Workbook:
    wb = Workbook()
    _invoices_sheet(wb, result.invoices, confidence_threshold)
    _tax_lines_sheet(wb, result.invoices)
    _vat_sheet(wb, result.invoices)
    _suppliers_sheet(wb, result.invoices)
    _issues_sheet(wb, result)
    return wb


def render_excel(result: Result, confidence_threshold: float = 0.7) -> bytes:
    buffer = io.BytesIO()
    build_workbook(result, confidence_threshold).save(buffer)
    return buffer.getvalue()


def write_excel(result: Result, target: Path, confidence_threshold: float = 0.7) -> Path:
    """Write the workbook atomically, so a viewer never sees a half-written file."""
    target.parent.mkdir(parents=True, exist_ok=True)
    fd, tmp_name = tempfile.mkstemp(dir=target.parent, suffix=".tmp")
    os.close(fd)
    tmp = Path(tmp_name)
    try:
        build_workbook(result, confidence_threshold).save(tmp)
        tmp.chmod(0o644)
        tmp.replace(target)
    finally:
        tmp.unlink(missing_ok=True)
    return target
