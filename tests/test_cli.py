import json
import threading
from pathlib import Path
from typing import Any

import pytest
from openpyxl import load_workbook
from typer.testing import CliRunner

from invoice_to_excel import __version__
from invoice_to_excel.cli import app
from invoice_to_excel.watch import FolderWatcher

runner = CliRunner()


def test_process_writes_excel_and_csv(invoice_folder: Path, tmp_path: Path) -> None:
    out = tmp_path / "output" / "invoices_report.xlsx"
    result = runner.invoke(app, ["process", str(invoice_folder), "--out", str(out)])
    assert result.exit_code == 0, result.output
    assert out.with_suffix(".csv").exists()
    assert "Summary: 18 invoices read (2 to review), 0 files with issues" in result.output

    wb = load_workbook(out)
    assert wb.sheetnames == ["Invoices", "Tax lines", "VAT 303", "Suppliers 347", "Issues"]
    ws = wb["Invoices"]
    statuses = [ws.cell(row=r, column=14).value for r in range(2, 20)]
    assert statuses.count("Review") == 2
    assert statuses.count("OK") == 16


def test_process_logs_one_structured_line_per_file(
    invoice_folder: Path, tmp_path: Path, caplog: pytest.LogCaptureFixture
) -> None:
    out = tmp_path / "r.xlsx"
    with caplog.at_level("INFO", logger="invoice_to_excel"):
        runner.invoke(app, ["process", str(invoice_folder), "--out", str(out)])
    processed = [r for r in caplog.records if r.getMessage() == "invoice processed"]
    assert len(processed) == 18
    assert {r.__dict__["status"] for r in processed} == {"ok", "review"}


def test_process_json_logs(invoice_folder: Path, tmp_path: Path) -> None:
    folder = tmp_path / "in"
    folder.mkdir()
    (folder / "x.pdf").write_bytes((invoice_folder / "supplies_2026-001.pdf").read_bytes())
    result = runner.invoke(
        app, ["--log-format", "json", "process", str(folder), "--out", str(tmp_path / "r.xlsx")]
    )
    assert result.exit_code == 0
    lines = [json.loads(line) for line in result.stderr.splitlines() if line.startswith("{")]
    assert any(
        entry["message"] == "invoice processed" and entry["file"] == "x.pdf" for entry in lines
    )


def test_process_missing_folder(tmp_path: Path) -> None:
    result = runner.invoke(app, ["process", str(tmp_path / "nothing")])
    assert result.exit_code == 2
    assert "does not exist" in result.stderr


def test_process_empty_folder(tmp_path: Path) -> None:
    result = runner.invoke(app, ["process", str(tmp_path), "--out", str(tmp_path / "r.xlsx")])
    assert result.exit_code == 1
    assert "No PDF found" in result.stderr


def test_csv_locale_and_no_csv(invoice_folder: Path, tmp_path: Path) -> None:
    es = tmp_path / "es.xlsx"
    assert (
        runner.invoke(
            app, ["process", str(invoice_folder), "--out", str(es), "--csv-locale", "es"]
        ).exit_code
        == 0
    )
    assert es.with_suffix(".csv").read_text(encoding="utf-8-sig").startswith("Supplier;Tax ID")
    plain = tmp_path / "plain.xlsx"
    assert (
        runner.invoke(
            app, ["process", str(invoice_folder), "--out", str(plain), "--no-csv"]
        ).exit_code
        == 0
    )
    assert not plain.with_suffix(".csv").exists()


def test_invalid_csv_locale(invoice_folder: Path, tmp_path: Path) -> None:
    result = runner.invoke(app, ["process", str(invoice_folder), "--csv-locale", "fr"])
    assert result.exit_code == 2


def test_no_ocr_leaves_scanned_pdfs_as_issues(invoice_folder: Path, tmp_path: Path) -> None:
    out = tmp_path / "r.xlsx"
    result = runner.invoke(app, ["process", str(invoice_folder), "--out", str(out), "--no-ocr"])
    assert "17 invoices read" in result.output
    assert "1 files with issues" in result.output
    assert load_workbook(out)["Issues"]["A2"].value == "renovations_scanned.pdf"


@pytest.mark.parametrize("mode", ["llm", "auto"])
def test_llm_modes_are_refused_without_opt_in(invoice_folder: Path, mode: str) -> None:
    result = runner.invoke(app, ["process", str(invoice_folder), "--extractor", mode])
    assert result.exit_code == 2
    assert "INVOICE_LLM_ENABLED" in result.stderr


def test_llm_mode_with_opt_in_but_no_key(
    invoice_folder: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    monkeypatch.setenv("INVOICE_LLM_ENABLED", "true")
    result = runner.invoke(app, ["process", str(invoice_folder), "--extractor", "llm"])
    assert result.exit_code == 2
    assert "ANTHROPIC_API_KEY" in result.stderr


def test_unknown_extractor_is_rejected(invoice_folder: Path) -> None:
    assert (
        runner.invoke(app, ["process", str(invoice_folder), "--extractor", "magic"]).exit_code == 2
    )


def test_version() -> None:
    result = runner.invoke(app, ["--version"])
    assert result.output.strip() == f"invoice-to-excel {__version__}"


def test_invalid_log_format() -> None:
    assert runner.invoke(app, ["--log-format", "xml", "process", "."]).exit_code == 2


def test_serve_passes_options_to_the_web_runner(monkeypatch: pytest.MonkeyPatch) -> None:
    calls: list[tuple[Any, ...]] = []
    import invoice_to_excel.web as web

    monkeypatch.setattr(web, "run", lambda settings, host, port: calls.append((host, port)))
    result = runner.invoke(app, ["serve", "--host", "0.0.0.0", "--port", "9000"])  # noqa: S104
    assert result.exit_code == 0
    assert calls == [("0.0.0.0", 9000)]  # noqa: S104


def test_serve_refuses_a_disabled_llm(monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setenv("INVOICE_EXTRACTOR", "auto")
    result = runner.invoke(app, ["serve"])
    assert result.exit_code == 2
    assert "INVOICE_LLM_ENABLED" in result.stderr


def test_watch_wires_the_watcher(
    invoice_folder: Path, tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    seen: dict[str, object] = {}

    def fake_run(self: FolderWatcher, stop: threading.Event, process_existing: bool = True) -> None:
        seen["folder"] = self.folder
        seen["existing"] = process_existing

    monkeypatch.setattr(FolderWatcher, "run", fake_run)
    inbox = tmp_path / "inbox"
    inbox.mkdir()
    out = tmp_path / "running.xlsx"
    result = runner.invoke(app, ["watch", str(inbox), "--out", str(out), "--no-existing"])
    assert result.exit_code == 0, result.output
    assert seen == {"folder": inbox.resolve(), "existing": False}
    assert f"Report: {out}" in result.output


def test_watch_missing_folder(tmp_path: Path) -> None:
    assert runner.invoke(app, ["watch", str(tmp_path / "nothing")]).exit_code == 2
