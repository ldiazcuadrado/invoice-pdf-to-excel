from datetime import date
from decimal import Decimal
from pathlib import Path

import pytest
from openpyxl import load_workbook

from invoice_to_excel.models import Invoice, Result, TaxLine
from invoice_to_excel.reporting import (
    SUPPLIER_347_THRESHOLD,
    quarterly_vat_summary,
    render_csv,
    render_excel,
    supplier_annual_totals,
    write_csv,
    write_excel,
)
from invoice_to_excel.reporting.tables import rate_label

D = Decimal


def vat_rows(
    result: Result,
) -> dict[tuple[str, Decimal | None], tuple[int, Decimal, Decimal, Decimal]]:
    return {
        (r.period, r.rate): (r.invoices, r.base, r.vat, r.surcharge)
        for r in quarterly_vat_summary(result.invoices)
    }


def test_quarterly_vat_groups_by_quarter_and_rate(result: Result) -> None:
    rows = vat_rows(result)
    # Q1: five invoices at 21%; the duplicate copy of 2026-014 is not counted twice.
    assert rows[("2026-T1", D(21))] == (5, D("6372.60"), D("1338.25"), D(0))
    assert rows[("2026-T2", D(4))] == (1, D("400.20"), D("16.01"), D(0))
    assert rows[("2026-T2", D(10))] == (1, D("1285.90"), D("128.59"), D(0))
    assert rows[("2026-T2", D(21))] == (5, D("7416.50"), D("1557.47"), D(0))


def test_multi_rate_and_surcharge_in_q3(result: Result) -> None:
    rows = vat_rows(result)
    assert rows[("2026-T3", D(4))] == (1, D("747.00"), D("29.88"), D(0))
    assert rows[("2026-T3", D(10))] == (3, D("1032.00"), D("103.20"), D("6.05"))
    assert rows[("2026-T3", D(21))] == (4, D("1252.00"), D("262.92"), D("32.86"))
    assert ("2026-T4", D(10)) in rows


def test_summary_is_ordered_by_period_then_rate_descending(result: Result) -> None:
    keys = [(r.period, r.rate) for r in quarterly_vat_summary(result.invoices)]
    assert keys == sorted(keys, key=lambda k: (k[0], -(k[1] or 0)))


def test_invoices_to_review_are_counted_in_their_bucket(result: Result) -> None:
    rows = {(r.period, r.rate): r for r in quarterly_vat_summary(result.invoices)}
    assert rows[("2026-T2", D(21))].to_review == 1  # supplies_2026-031 has a wrong total
    assert rows[("2026-T1", D(21))].to_review == 0


def test_undated_invoice_goes_to_its_own_bucket() -> None:
    invoice = Invoice(filename="a.pdf", tax_lines=[TaxLine(base=D(10), rate=D(21), vat=D("2.10"))])
    (row,) = quarterly_vat_summary([invoice])
    assert row.period == "No date"


def test_supplier_totals_flag_the_347_threshold(result: Result) -> None:
    rows = {r.supplier: r for r in supplier_annual_totals(result.invoices)}
    assert D("3005.06") == SUPPLIER_347_THRESHOLD
    marta = rows["Marta Ficticia Ruiz"]
    assert (marta.year, marta.invoices, marta.gross) == (2026, 2, D("4779.50"))
    assert marta.over_threshold
    assert rows["Talleres Inventados, S.L."].over_threshold
    assert rows["Reformas Escaneadas, S.L."].over_threshold
    suministros = rows["Suministros Ficticios, S.L."]
    assert suministros.invoices == 4  # duplicate excluded
    assert suministros.gross == D("2031.11")  # base + VAT of the parts, not the wrong stated total
    assert not suministros.over_threshold
    assert not rows["Consultoría Falsa, S.L.P."].over_threshold


def test_irpf_does_not_reduce_the_347_amount() -> None:
    invoice = Invoice(
        filename="a.pdf",
        supplier="Pro",
        tax_id="12345678Z",
        invoice_date=date(2026, 5, 1),
        tax_lines=[TaxLine(base=D(3000), rate=D(21), vat=D(630))],
        irpf=D(450),
        total=D(3180),
    )
    (row,) = supplier_annual_totals([invoice])
    assert row.gross == D(3630)
    assert row.over_threshold


def test_threshold_is_exclusive() -> None:
    invoice = Invoice(
        filename="a.pdf",
        supplier="S",
        tax_id="B87654323",
        invoice_date=date(2026, 1, 1),
        tax_lines=[TaxLine(base=D("3005.06"), rate=D(0), vat=D(0))],
    )
    assert not supplier_annual_totals([invoice])[0].over_threshold
    assert supplier_annual_totals([invoice], threshold=D("3005.05"))[0].over_threshold


def test_same_supplier_in_two_years_is_two_rows() -> None:
    def invoice(name: str, year: int) -> Invoice:
        return Invoice(
            filename=name,
            supplier="S",
            tax_id="B87654323",
            invoice_date=date(year, 6, 1),
            tax_lines=[TaxLine(base=D(100), rate=D(21), vat=D(21))],
        )

    rows = supplier_annual_totals([invoice("a", 2025), invoice("b", 2026)])
    assert [r.year for r in rows] == [2025, 2026]


def test_rate_label(by_file: dict[str, Invoice]) -> None:
    assert rate_label(by_file["supplies_2026-001.pdf"]) == "21%"
    assert rate_label(by_file["hardware_F-2026-0118.pdf"]) == "21% + 10%"
    assert rate_label(by_file["wholesale_B-5521.pdf"], decimal_comma=True) == "21% + 10%"
    assert rate_label(Invoice(filename="x")) == ""


# -- Excel -----------------------------------------------------------------------


def test_workbook_layout(result: Result, tmp_path: Path) -> None:
    out = write_excel(result, tmp_path / "out" / "report.xlsx")
    wb = load_workbook(out)
    assert wb.sheetnames == ["Invoices", "Tax lines", "VAT 303", "Suppliers 347", "Issues"]
    ws = wb["Invoices"]
    assert ws.freeze_panes == "A2"
    assert ws.auto_filter.ref == f"A1:O{len(result.invoices) + 1}"
    assert [c.value for c in ws[1]][:6] == [
        "Supplier", "Tax ID", "Invoice no.", "Date", "Net amount", "VAT %",
    ]  # fmt: skip
    assert wb["Issues"]["A2"].value == "No issues: every PDF could be read."
    assert not list(out.parent.glob("*.tmp"))


def test_invoices_sheet_shows_the_vat_breakdown(result: Result, tmp_path: Path) -> None:
    ws = load_workbook(write_excel(result, tmp_path / "r.xlsx"))["Invoices"]
    rows = {ws.cell(row=r, column=13).value: r for r in range(2, ws.max_row + 1)}
    row = rows["hardware_F-2026-0118.pdf"]
    assert ws.cell(row=row, column=6).value == "21% + 10%"
    assert ws.cell(row=row, column=7).value == pytest.approx(100.38)
    assert ws.cell(row=rows["wholesale_B-5521.pdf"], column=8).value == pytest.approx(21.13)


def test_tax_lines_sheet_has_one_row_per_rate(result: Result, tmp_path: Path) -> None:
    ws = load_workbook(write_excel(result, tmp_path / "r.xlsx"))["Tax lines"]
    expected = sum(len(i.tax_lines) for i in result.invoices)
    assert ws.max_row == expected + 1
    rows = [[c.value for c in row] for row in ws.iter_rows(min_row=2)]
    assert any(r[1] == "2026/0077" and r[3] == "4%" and r[4] == pytest.approx(747.0) for r in rows)


def test_review_rows_and_low_confidence_cells_are_highlighted(
    result: Result, tmp_path: Path
) -> None:
    ws = load_workbook(write_excel(result, tmp_path / "r.xlsx"))["Invoices"]
    rows = {ws.cell(row=r, column=13).value: r for r in range(2, ws.max_row + 1)}
    wrong_total = rows["supplies_2026-031.pdf"]
    assert ws.cell(row=wrong_total, column=14).value == "Review"
    assert ws.cell(row=wrong_total, column=1).fill.fgColor.rgb.endswith("FFC7CE")
    assert ws.cell(row=wrong_total, column=10).fill.fgColor.rgb.endswith("FFD966")  # total
    comment = ws.cell(row=wrong_total, column=10).comment
    assert comment is not None
    assert "Confidence" in comment.text
    marta = rows["marta_design_03.pdf"]  # supplier name is a weak guess
    assert ws.cell(row=marta, column=1).fill.fgColor.rgb.endswith("FFD966")
    assert ws.cell(row=marta, column=2).fill.fgColor.rgb in ("00000000", "0")
    clean = rows["supplies_2026-001.pdf"]
    assert ws.cell(row=clean, column=1).fill.fill_type is None


def test_vat_sheet_has_subtotals_and_grand_total(result: Result, tmp_path: Path) -> None:
    ws = load_workbook(write_excel(result, tmp_path / "r.xlsx"))["VAT 303"]
    cells = {
        (r, c): ws.cell(row=r, column=c).value
        for r in range(1, ws.max_row + 1)
        for c in range(1, 8)
    }
    assert ws["A1"].value == "Quarter"
    labels = [v for (r, c), v in cells.items() if c == 1 and isinstance(v, str)]
    assert "2026-T1 total" in labels
    assert "ALL" in labels
    t1 = next(r for (r, c), v in cells.items() if c == 1 and v == "2026-T1 total")
    assert ws.cell(row=t1, column=5).value == pytest.approx(6372.60)
    assert ws.cell(row=t1, column=6).value == pytest.approx(1338.25)


def test_347_sheet_flags_suppliers_over_the_threshold(result: Result, tmp_path: Path) -> None:
    ws = load_workbook(write_excel(result, tmp_path / "r.xlsx"))["Suppliers 347"]
    flagged = {
        ws.cell(row=r, column=2).value
        for r in range(2, ws.max_row + 1)
        if ws.cell(row=r, column=6).value == "YES"
    }
    assert "Marta Ficticia Ruiz" in flagged
    assert "Consultoría Falsa, S.L.P." not in flagged


def test_issues_sheet_lists_file_issues(tmp_path: Path) -> None:
    from invoice_to_excel.models import FileIssue

    result = Result(file_issues=[FileIssue(filename="bad.pdf", reason="broken")])
    ws = load_workbook(write_excel(result, tmp_path / "r.xlsx"))["Issues"]
    assert [ws["A2"].value, ws["B2"].value] == ["bad.pdf", "broken"]


def test_empty_result_still_produces_a_workbook(tmp_path: Path) -> None:
    wb = load_workbook(write_excel(Result(), tmp_path / "r.xlsx"))
    assert wb["Invoices"].max_row == 1


def test_render_excel_returns_xlsx_bytes(result: Result) -> None:
    assert render_excel(result)[:2] == b"PK"


# -- CSV -------------------------------------------------------------------------


def test_csv_locales(result: Result, tmp_path: Path) -> None:
    en = write_csv(result, tmp_path / "en.csv").read_text(encoding="utf-8-sig").splitlines()
    es = write_csv(result, tmp_path / "es.csv", "es").read_text(encoding="utf-8-sig").splitlines()
    assert en[0].startswith("Supplier,Tax ID,Invoice no.,Date,")
    assert es[0].startswith("Supplier;Tax ID;Invoice no.;Date;")
    assert "252.60" in "\n".join(en)
    assert "2026-01-15" in "\n".join(en)
    assert "252,60" in "\n".join(es)
    assert "15/01/2026" in "\n".join(es)
    assert len(en) == len(result.invoices) + 1


def test_csv_contains_the_breakdown_and_status(result: Result) -> None:
    text = render_csv(result)
    assert "21% + 10% + 4%" in text
    assert "Possible duplicate invoice" in text


def test_csv_rejects_unknown_locale(result: Result) -> None:
    with pytest.raises(ValueError, match="locale"):
        render_csv(result, "fr")
