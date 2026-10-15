import shutil
import threading
import time
from collections.abc import Callable
from pathlib import Path

import pytest
from openpyxl import load_workbook

from invoice_to_excel.config import Settings
from invoice_to_excel.extraction import Pipeline, build_pipeline
from invoice_to_excel.models import FileIssue, Invoice
from invoice_to_excel.watch import FolderWatcher, RunningReport, _move_unique

GOOD = "supplies_2026-001.pdf"


@pytest.fixture
def inbox(tmp_path: Path) -> Path:
    folder = tmp_path / "inbox"
    folder.mkdir()
    return folder


@pytest.fixture
def report(tmp_path: Path) -> RunningReport:
    return RunningReport(tmp_path / "reports" / "running.xlsx")


def make_watcher(
    inbox: Path,
    pipeline: Pipeline,
    report: RunningReport,
    settle_seconds: float = 1.0,
    clock: Callable[[], float] = time.monotonic,
) -> FolderWatcher:
    return FolderWatcher(inbox, pipeline, report, settle_seconds=settle_seconds, clock=clock)


def wait_for(condition: Callable[[], bool], timeout: float = 15.0) -> bool:
    deadline = time.monotonic() + timeout
    while time.monotonic() < deadline:
        if condition():
            return True
        time.sleep(0.05)
    return False


def test_handle_moves_invoice_to_processed(
    inbox: Path, invoice_folder: Path, pipeline: Pipeline, report: RunningReport
) -> None:
    shutil.copy(invoice_folder / GOOD, inbox)
    watcher = make_watcher(inbox, pipeline, report)
    outcome = watcher.handle(inbox / GOOD)
    assert isinstance(outcome, Invoice)
    assert (inbox / "processed" / GOOD).exists()
    assert not (inbox / GOOD).exists()
    assert len(report.result.invoices) == 1
    assert report.xlsx.exists()
    assert report.xlsx.with_suffix(".csv").exists()


def test_handle_moves_unreadable_file_to_failed(
    inbox: Path, pipeline: Pipeline, report: RunningReport
) -> None:
    (inbox / "broken.pdf").write_bytes(b"not a pdf")
    outcome = make_watcher(inbox, pipeline, report).handle(inbox / "broken.pdf")
    assert isinstance(outcome, FileIssue)
    assert (inbox / "failed" / "broken.pdf").exists()
    assert load_workbook(report.xlsx)["Issues"]["A2"].value == "broken.pdf"


def test_report_grows_and_flags_duplicates_across_files(
    inbox: Path, invoice_folder: Path, pipeline: Pipeline, report: RunningReport
) -> None:
    watcher = make_watcher(inbox, pipeline, report)
    for name in (GOOD, "supplies_2026-014.pdf", "supplies_2026-014_copy.pdf"):
        shutil.copy(invoice_folder / name, inbox)
        watcher.handle(inbox / name)
        rows = load_workbook(report.xlsx)["Invoices"].max_row
    assert rows >= 4
    assert [i.duplicate_of for i in report.result.invoices] == [None, None, "supplies_2026-014.pdf"]


def test_running_report_survives_restarts(
    inbox: Path, invoice_folder: Path, pipeline: Pipeline, report: RunningReport
) -> None:
    shutil.copy(invoice_folder / GOOD, inbox)
    make_watcher(inbox, pipeline, report).handle(inbox / GOOD)
    (inbox / "broken.pdf").write_bytes(b"x")
    make_watcher(inbox, pipeline, report).handle(inbox / "broken.pdf")

    reloaded = RunningReport(report.xlsx)
    assert [i.filename for i in reloaded.result.invoices] == [GOOD]
    assert [i.filename for i in reloaded.result.file_issues] == ["broken.pdf"]
    shutil.copy(invoice_folder / "supplies_2026-014.pdf", inbox)
    make_watcher(inbox, pipeline, reloaded).handle(inbox / "supplies_2026-014.pdf")
    assert load_workbook(report.xlsx)["Invoices"].max_row >= 3


def test_blank_ledger_lines_are_ignored(tmp_path: Path) -> None:
    xlsx = tmp_path / "r.xlsx"
    xlsx.with_suffix(".jsonl").write_text("\n\n", encoding="utf-8")
    assert RunningReport(xlsx).result.invoices == []


def test_csv_can_be_disabled(tmp_path: Path) -> None:
    report = RunningReport(tmp_path / "r.xlsx", write_csv_file=False)
    report.write()
    assert report.xlsx.exists()
    assert not report.xlsx.with_suffix(".csv").exists()


def test_name_collisions_do_not_overwrite(tmp_path: Path) -> None:
    target = tmp_path / "done"
    names = []
    for _ in range(3):
        source = tmp_path / "a.pdf"
        source.write_bytes(b"x")
        names.append(_move_unique(source, target).name)
    assert names == ["a.pdf", "a.1.pdf", "a.2.pdf"]


def run_in_background(
    watcher: FolderWatcher, **kw: bool
) -> tuple[threading.Event, threading.Thread]:
    stop = threading.Event()
    thread = threading.Thread(target=watcher.run, args=(stop,), kwargs=kw, daemon=True)
    thread.start()
    return stop, thread


def test_watcher_picks_up_new_files(
    inbox: Path, invoice_folder: Path, pipeline: Pipeline, report: RunningReport
) -> None:
    watcher = make_watcher(inbox, pipeline, report, settle_seconds=0.2)
    stop, thread = run_in_background(watcher)
    try:
        assert wait_for(watcher.processed_dir.exists)
        shutil.copy(invoice_folder / GOOD, inbox / "incoming.pdf")
        (inbox / "notes.txt").write_text("ignored")
        assert wait_for(lambda: (inbox / "processed" / "incoming.pdf").exists())
    finally:
        stop.set()
        thread.join(timeout=10)
    assert not thread.is_alive()
    assert len(report.result.invoices) == 1
    assert (inbox / "notes.txt").exists()


def test_watcher_processes_existing_files_and_sorts_failures(
    inbox: Path, invoice_folder: Path, pipeline: Pipeline, report: RunningReport
) -> None:
    shutil.copy(invoice_folder / GOOD, inbox)
    (inbox / "broken.pdf").write_bytes(b"nope")
    watcher = make_watcher(inbox, pipeline, report, settle_seconds=0.2)
    stop, thread = run_in_background(watcher)
    try:
        assert wait_for(lambda: (inbox / "processed" / GOOD).exists())
        assert wait_for(lambda: (inbox / "failed" / "broken.pdf").exists())
    finally:
        stop.set()
        thread.join(timeout=10)
    assert len(report.result.invoices) == 1
    assert len(report.result.file_issues) == 1


def test_watcher_can_skip_existing_files(
    inbox: Path, invoice_folder: Path, pipeline: Pipeline, report: RunningReport
) -> None:
    shutil.copy(invoice_folder / GOOD, inbox)
    watcher = make_watcher(inbox, pipeline, report, settle_seconds=0.1)
    stop, thread = run_in_background(watcher, process_existing=False)
    try:
        time.sleep(1.0)
    finally:
        stop.set()
        thread.join(timeout=10)
    assert (inbox / GOOD).exists()
    assert report.result.invoices == []


def test_file_is_not_processed_while_it_is_still_growing(
    inbox: Path, invoice_folder: Path, pipeline: Pipeline, report: RunningReport
) -> None:
    now = [0.0]
    watcher = make_watcher(inbox, pipeline, report, settle_seconds=5, clock=lambda: now[0])
    path = inbox / "slow.pdf"
    data = (invoice_folder / GOOD).read_bytes()
    path.write_bytes(data[:100])
    pending: dict[Path, tuple[int, float]] = {path: (-1, 0.0)}

    watcher._process_stable(pending)
    now[0] = 10
    path.write_bytes(data[:200])  # grew since the last look: the clock restarts
    watcher._process_stable(pending)
    assert path in pending
    assert path.exists()

    path.write_bytes(data)
    watcher._process_stable(pending)  # size changed again
    now[0] = 20
    watcher._process_stable(pending)  # stable for >= 5 s
    assert path not in pending
    assert (inbox / "processed" / "slow.pdf").exists()


def test_vanished_files_are_dropped(inbox: Path, pipeline: Pipeline, report: RunningReport) -> None:
    watcher = make_watcher(inbox, pipeline, report)
    pending = {inbox / "gone.pdf": (-1, 0.0)}
    watcher._process_stable(pending)
    assert pending == {}


def test_errors_while_filing_do_not_stop_the_watcher(
    inbox: Path, invoice_folder: Path, pipeline: Pipeline, report: RunningReport,
    monkeypatch: pytest.MonkeyPatch,
) -> None:  # fmt: skip
    watcher = make_watcher(inbox, pipeline, report, settle_seconds=0)

    def boom(path: Path) -> Invoice | FileIssue:
        raise OSError("disk full")

    monkeypatch.setattr(watcher, "handle", boom)
    shutil.copy(invoice_folder / GOOD, inbox)
    pending = {inbox / GOOD: (inbox.joinpath(GOOD).stat().st_size, 0.0)}
    watcher._process_stable(pending)
    assert pending == {}


def test_pipeline_fixture_is_rules_only(settings: Settings) -> None:
    assert build_pipeline(settings).mode == "rules"
