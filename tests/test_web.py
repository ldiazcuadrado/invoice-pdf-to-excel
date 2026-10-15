import csv
import io
import tempfile
from pathlib import Path
from typing import Any

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from invoice_to_excel.config import Settings
from invoice_to_excel.web import create_app

from .fakes import GOOD_PAYLOAD, FakeClient


@pytest.fixture
def client() -> TestClient:
    return TestClient(create_app(Settings(_env_file=None)))


def upload(invoice_folder: Path, *names: str) -> list[tuple[str, tuple[str, bytes, str]]]:
    return [
        ("files", (name, (invoice_folder / name).read_bytes(), "application/pdf")) for name in names
    ]


def test_index_page_and_static_assets(client: TestClient) -> None:
    page = client.get("/")
    assert page.status_code == 200
    assert "Invoice PDF to Excel" in page.text
    assert client.get("/static/app.js").status_code == 200
    assert client.get("/static/style.css").status_code == 200


def test_security_headers(client: TestClient) -> None:
    response = client.get("/")
    assert "default-src 'self'" in response.headers["content-security-policy"]
    assert response.headers["x-content-type-options"] == "nosniff"
    assert response.headers["cache-control"] == "no-store"


def test_health_and_config(client: TestClient) -> None:
    assert client.get("/healthz").json() == {"status": "ok"}
    config = client.get("/api/config").json()
    assert config["max_upload_mb"] == 10
    assert config["sends_data_to_third_party"] is False


def test_config_reports_when_data_leaves_the_machine() -> None:
    settings = Settings(_env_file=None, llm_enabled=True)
    from invoice_to_excel.extraction import build_pipeline

    pipeline = build_pipeline(settings, "auto", llm_client=FakeClient(GOOD_PAYLOAD))
    app_client = TestClient(create_app(settings, pipeline))
    assert app_client.get("/api/config").json()["sends_data_to_third_party"] is True


def test_process_returns_extracted_invoices(client: TestClient, invoice_folder: Path) -> None:
    response = client.post(
        "/api/process",
        files=upload(invoice_folder, "supplies_2026-001.pdf", "hardware_F-2026-0118.pdf"),
    )
    assert response.status_code == 200
    body = response.json()
    assert [i["filename"] for i in body["invoices"]] == [
        "supplies_2026-001.pdf",
        "hardware_F-2026-0118.pdf",
    ]
    first, second = body["invoices"]
    assert first["supplier"] == "Suministros Ficticios, S.L."
    assert first["total"] == "305.65"
    assert first["is_ok"] is True
    assert first["overall_confidence"] > 0.9
    assert len(second["tax_lines"]) == 2
    assert body["file_issues"] == []


def test_process_flags_duplicates_within_a_request(
    client: TestClient, invoice_folder: Path
) -> None:
    files = upload(invoice_folder, "supplies_2026-014.pdf", "supplies_2026-014_copy.pdf")
    body = client.post("/api/process", files=files).json()
    assert body["invoices"][1]["duplicate_of"] == "supplies_2026-014.pdf"


def test_process_reports_unreadable_and_non_pdf_files(client: TestClient) -> None:
    files = [
        ("files", ("notes.pdf", b"plain text", "application/pdf")),
        ("files", ("broken.pdf", b"%PDF-1.4 garbage", "application/pdf")),
    ]
    body = client.post("/api/process", files=files).json()
    assert body["invoices"] == []
    reasons = {i["filename"]: i["reason"] for i in body["file_issues"]}
    assert reasons["notes.pdf"] == "Not a PDF file"
    assert "Could not open" in reasons["broken.pdf"]


def test_filenames_cannot_escape_the_upload_directory(
    client: TestClient, invoice_folder: Path
) -> None:
    data = (invoice_folder / "supplies_2026-001.pdf").read_bytes()
    files = [("files", ("../../etc/evil.pdf", data, "application/pdf"))]
    body = client.post("/api/process", files=files).json()
    assert body["invoices"][0]["filename"] == "evil.pdf"


def test_file_without_name_or_extension_is_handled(
    client: TestClient, invoice_folder: Path
) -> None:
    data = (invoice_folder / "supplies_2026-001.pdf").read_bytes()
    body = client.post(
        "/api/process", files=[("files", ("invoice", data, "application/pdf"))]
    ).json()
    assert body["invoices"][0]["filename"] == "invoice.pdf"


def test_oversized_upload_is_rejected() -> None:
    app_client = TestClient(create_app(Settings(_env_file=None, web_max_upload_mb=1)))
    payload = b"%PDF-" + b"0" * (1024 * 1024)
    response = app_client.post(
        "/api/process", files=[("files", ("big.pdf", payload, "application/pdf"))]
    )
    assert response.status_code == 413
    assert "1 MB" in response.json()["detail"]


def test_too_many_files_are_rejected() -> None:
    app_client = TestClient(create_app(Settings(_env_file=None, web_max_files=2)))
    files = [("files", (f"{i}.pdf", b"%PDF-1.4", "application/pdf")) for i in range(3)]
    assert app_client.post("/api/process", files=files).status_code == 413


def test_uploads_are_deleted_after_processing(
    client: TestClient, invoice_folder: Path, tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    scratch = tmp_path / "scratch"
    scratch.mkdir()
    monkeypatch.setattr(tempfile, "tempdir", str(scratch))
    assert (
        client.post(
            "/api/process", files=upload(invoice_folder, "supplies_2026-001.pdf")
        ).status_code
        == 200
    )
    assert list(scratch.iterdir()) == []


def test_uploads_are_deleted_when_the_request_fails(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch, invoice_folder: Path
) -> None:
    scratch = tmp_path / "scratch"
    scratch.mkdir()
    monkeypatch.setattr(tempfile, "tempdir", str(scratch))
    app_client = TestClient(create_app(Settings(_env_file=None, web_max_upload_mb=1)))
    big = b"%PDF-" + b"0" * (1024 * 1024)
    files = [
        (
            "files",
            ("ok.pdf", (invoice_folder / "supplies_2026-001.pdf").read_bytes(), "application/pdf"),
        ),
        ("files", ("big.pdf", big, "application/pdf")),
    ]
    assert app_client.post("/api/process", files=files).status_code == 413
    assert list(scratch.iterdir()) == []


def processed(client: TestClient, invoice_folder: Path, *names: str) -> dict[str, Any]:
    body: dict[str, Any] = client.post("/api/process", files=upload(invoice_folder, *names)).json()
    return body


def test_report_adds_accountant_summaries(client: TestClient, invoice_folder: Path) -> None:
    result = processed(
        client,
        invoice_folder,
        "marta_design_03.pdf",
        "marta_design_07.pdf",
        "hardware_F-2026-0118.pdf",
    )
    report = client.post("/api/report", json=result).json()
    assert {r["period"] for r in report["vat_quarterly"]} == {"2026-T1", "2026-T2", "2026-T3"}
    marta = next(r for r in report["suppliers_347"] if r["supplier"] == "Marta Ficticia Ruiz")
    assert marta["over_threshold"] is True
    assert report["confidence_threshold"] == 0.7
    assert len(report["invoices"]) == 3


def test_report_flags_duplicates_across_separate_uploads(
    client: TestClient, invoice_folder: Path
) -> None:
    first = processed(client, invoice_folder, "supplies_2026-014.pdf")
    second = processed(client, invoice_folder, "supplies_2026-014_copy.pdf")
    merged = {
        "invoices": first["invoices"] + second["invoices"],
        "file_issues": first["file_issues"] + second["file_issues"],
    }
    report = client.post("/api/report", json=merged).json()
    assert [i["duplicate_of"] for i in report["invoices"]] == [None, "supplies_2026-014.pdf"]
    assert sum(r["invoices"] for r in report["vat_quarterly"]) == 1


def test_report_rejects_malformed_bodies(client: TestClient) -> None:
    assert client.post("/api/report", json={"invoices": [{"nope": 1}]}).status_code == 422


def test_report_size_limit() -> None:
    from invoice_to_excel.web import app as web_app

    with pytest.MonkeyPatch.context() as mp:
        mp.setattr(web_app, "MAX_REPORT_INVOICES", 1)
        app_client = TestClient(create_app(Settings(_env_file=None)))
        invoice = {"filename": "a.pdf"}
        assert (
            app_client.post("/api/report", json={"invoices": [invoice, invoice]}).status_code == 413
        )


def test_export_xlsx(client: TestClient, invoice_folder: Path) -> None:
    result = processed(client, invoice_folder, "supplies_2026-001.pdf", "hardware_F-2026-0118.pdf")
    response = client.post("/api/export/xlsx", json=result)
    assert response.status_code == 200
    assert "spreadsheetml" in response.headers["content-type"]
    assert 'filename="invoices_report.xlsx"' in response.headers["content-disposition"]
    wb = load_workbook(io.BytesIO(response.content))
    assert wb["Invoices"]["A2"].value == "Suministros Ficticios, S.L."
    assert wb["Tax lines"].max_row == 4


def test_export_csv(client: TestClient, invoice_folder: Path) -> None:
    result = processed(client, invoice_folder, "supplies_2026-001.pdf")
    response = client.post("/api/export/csv", json=result)
    assert response.status_code == 200
    assert response.content.startswith(b"\xef\xbb\xbf")
    rows = list(csv.reader(io.StringIO(response.content.decode("utf-8-sig"))))
    assert rows[1][2] == "2026-001"
    assert rows[1][9] == "305.65"


def test_export_rejects_unknown_format(client: TestClient) -> None:
    assert client.post("/api/export/pdf", json={}).status_code == 422


def test_nothing_is_persisted_by_the_server(
    client: TestClient, invoice_folder: Path, tmp_path: Path
) -> None:
    before = set(Path.cwd().rglob("*"))
    result = processed(client, invoice_folder, "supplies_2026-001.pdf")
    client.post("/api/report", json=result)
    client.post("/api/export/xlsx", json=result)
    assert set(Path.cwd().rglob("*")) == before
